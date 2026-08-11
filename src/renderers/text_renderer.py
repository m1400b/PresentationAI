"""
PresentationAI

Text Renderer
"""

from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from src.models.elements.text_element import TextElement


class TextRenderer:
    """
    Renders a TextElement into a PowerPoint textbox.
    """

    PX_TO_CM = 37.7952755906

    # -------------------------------------------------

    @classmethod
    def px(cls, value: float):

        return Cm(value / cls.PX_TO_CM)

    # -------------------------------------------------

    def render(
        self,
        ppt_slide,
        element: TextElement,
    ):

        textbox = ppt_slide.shapes.add_textbox(

            self.px(element.x),

            self.px(element.y),

            self.px(element.width),

            self.px(element.height),

        )

        frame = textbox.text_frame

        frame.clear()

        frame.word_wrap = True

        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0

        paragraph = frame.paragraphs[0]

        paragraph.text = element.text

        # ---------------------------------------------
        # Alignment
        # ---------------------------------------------

        alignment = (element.alignment or "").lower()

        if alignment == "center":

            paragraph.alignment = PP_ALIGN.CENTER

        elif alignment == "right":

            paragraph.alignment = PP_ALIGN.RIGHT

        else:

            paragraph.alignment = PP_ALIGN.LEFT

        # ---------------------------------------------
        # Font
        # ---------------------------------------------

        font = paragraph.font

        font.name = element.font_family

        font.size = Pt(element.font_size)

        font.bold = element.bold

        font.italic = element.italic

        font.underline = element.underline

        # ---------------------------------------------
        # Color
        # ---------------------------------------------

        try:

            color = element.color.replace("#", "")

            if len(color) == 6:

                font.color.rgb = RGBColor.from_string(color)

        except Exception:

            pass